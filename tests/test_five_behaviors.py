"""Real files/real Tesseract, real ingestion/retrieval/reasoning. No mocked OCR.

The corpus is synthetic, isolated from the institutional database and contains
known facts plus counterexamples. No provider is necessary for these contracts.
"""
import asyncio
import io
import re
from pathlib import Path

import pytest
from fastapi import HTTPException
from docx import Document
from openpyxl import Workbook
import pymupdf
from PIL import Image, ImageDraw, ImageFont

from api import ingestion
from api.auth import require_admin
from api.context_engine import verify_answer
from api.document_pages import extract_pages, pages_ready
from api.expansion import record_document_pipeline, pipeline_documents
from api.ingestion import extract_text, ingest_module
from api.orchestration import answer, _format_external_assist_answer
from api.policies import policy_for
from api.relational_reasoning import analyze, render
from api.retrieval import retrieve, warm_module_index, _index, _normalized_index
from api.structured_data import analyze_structured_question, load_structured_document


@pytest.fixture
def corpus(tmp_path, monkeypatch):
    monkeypatch.setenv("SOFIA_STORAGE_MODE", "developer")
    for key in ("SOFIA_POSTGRES_URL", "DATABASE_URL", "OPENAI_API_KEY", "GEMINI_API_KEY", "ANTHROPIC_API_KEY", "SOFIA_ENCRYPTION_KEY"):
        monkeypatch.delenv(key, raising=False)
    monkeypatch.setenv("SOFIA_EMBEDDINGS_QUERY_ENABLED", "false")
    monkeypatch.setattr(ingestion, "CACHE_ROOT", tmp_path / "cache")
    root = tmp_path / "knowledge"
    (root / "infraestrutura").mkdir(parents=True)
    return root


def source(root, name, text, module="infraestrutura"):
    path = root / module / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return path


def prepare(root, *paths):
    for path in paths:
        result = record_document_pipeline(root, path.parent.name, path)
        assert result["status"] == "READY", result
    warm_module_index(root, paths[0].parent.name, force=True)


def ask(root, question, module="infraestrutura"):
    return asyncio.run(answer(root=root, module_id=module, provider="auto", question=question, history=[], external_allowed=False))


def scanned_pdf(path, text, mixed=False):
    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 36)
    draw.multiline_text((90, 150), text, fill="black", font=font, spacing=18)
    buffer = io.BytesIO(); image.save(buffer, format="PNG")
    with pymupdf.open() as doc:
        if mixed:
            page = doc.new_page()
            page.insert_text((60, 90), "Manual Atlas: a rotina de backup deve ser registrada diariamente.", fontsize=12)
        page = doc.new_page()
        page.insert_image(page.rect, stream=buffer.getvalue())
        doc.save(path)


@pytest.mark.parametrize("suffix", ["txt", "xml", "docx", "csv", "xlsx", "pdf"])
def test_formats_extract_and_retrieve_actual_document(corpus, suffix):
    path = corpus / "infraestrutura" / f"manual_atlas.{suffix}"
    phrase = "O servidor Atlas exige backup diario e retencao de 30 dias."
    if suffix == "docx":
        doc = Document(); doc.add_paragraph("Manual operacional Atlas.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Servidor"; table.cell(0, 1).text = "Regra"
        table.cell(1, 0).text = "Atlas"; table.cell(1, 1).text = phrase
        doc.save(path)
    elif suffix == "xlsx":
        workbook = Workbook(); workbook.active.title = "Politicas"
        workbook.active.append(["Servidor", "Regra"]); workbook.active.append(["Atlas", phrase]); workbook.save(path)
    elif suffix == "pdf":
        with pymupdf.open() as doc:
            page = doc.new_page(); page.insert_text((60, 90), phrase); doc.save(path)
    else:
        path.write_text({"txt": phrase, "xml": f'<politica sistema="Atlas"><regra>{phrase}</regra></politica>', "csv": f'Servidor;Regra\nAtlas;{phrase}\n'}[suffix], encoding="utf-8")
    prepare(corpus, path)
    result = retrieve(corpus, "infraestrutura", "Qual a regra de backup do arquivo manual_atlas?", policy_for("infraestrutura"))
    assert result.has_quality_evidence
    assert "30 dias" in result.context
    assert path.name in result.sources
    if suffix == "pdf":
        assert result.evidence[0].chunk.page == 1


def test_summary_and_inference_across_documents(corpus):
    rule = source(corpus, "politica.txt", "Todo servidor de produção exige backup diário. O prazo de retenção dos logs é de 30 dias.")
    asset = source(corpus, "inventario.txt", "Atlas é um servidor de produção. O responsável pela operação do Atlas é a equipe de infraestrutura.")
    prepare(corpus, rule, asset)
    summary = ask(corpus, "Resuma o arquivo politica")
    assert "backup diário" in summary.answer
    assert "30 dias" in summary.answer
    assert "[E" in summary.answer
    assert summary.verification_status == "verified"
    result = ask(corpus, "Compare politica e inventario e faça uma inferência sobre Atlas")
    assert "Inferência" in result.answer
    assert "Atlas exige backup diário" in result.answer
    assert set(result.sources) == {rule.name, asset.name}
    assert result.verification_status == "verified"
    inference = next(r for r in result.context_package["relational_analysis"]["relations"] if r["inferred"])
    assert len(inference["premises"]) == 2


def test_conflicts_and_no_invented_causality(corpus):
    first = source(corpus, "politica.txt", "O prazo de retenção dos logs é de 30 dias. A sobrecarga está correlacionada com latência elevada.")
    second = source(corpus, "auditoria.txt", "O prazo de retenção dos logs é de 90 dias. A sobrecarga pode causar interrupções no serviço.")
    prepare(corpus, first, second)
    result = ask(corpus, "Compare politica e auditoria: conflitos, correlação e causalidade")
    assert result.verification_status == "verified"
    assert "divergência textual" in result.answer
    assert "30 dias" in result.answer and "90 dias" in result.answer
    assert "Correlação documentada" in result.answer
    assert "Hipótese documentada" in result.answer
    assert "não demonstra causalidade" in result.answer


def test_subclasses_and_patterns(corpus):
    first = source(corpus, "taxonomia.txt", "Um servidor crítico é um tipo de servidor de produção. Todo servidor de produção exige backup diário.")
    second = source(corpus, "inventario.txt", "Atlas é um servidor crítico. Todo servidor de produção exige backup diário.")
    prepare(corpus, first, second)
    result = ask(corpus, "Compare taxonomia e inventario: subclasses, padrões e inferências")
    relations = result.context_package["relational_analysis"]["relations"]
    assert any(r["type"] == "subclass" and not r["inferred"] for r in relations)
    assert "Atlas exige backup diário" in result.answer
    assert any(p["kind"] == "pattern" for p in result.context_package["relational_analysis"]["comparisons"])


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_typed_numeric_aggregation_and_filter(corpus, suffix):
    path = corpus / "infraestrutura" / f"custos.{suffix}"
    if suffix == "csv":
        path.write_text("Equipe;Valor\nTI;1.234,50\nTI;2.000,25\nRH;500,00\nTI;\n", encoding="utf-8")
    else:
        book = Workbook(); book.active.title = "Custos"
        for row in [["Equipe", "Valor"], ["TI", 1234.5], ["TI", 2000.25], ["RH", 500], ["TI", None]]:
            book.active.append(row)
        book.save(path)
    prepare(corpus, path)
    result = ask(corpus, "Qual a soma de Valor para TI no arquivo custos?")
    data = result.context_package["structured_data"]
    assert data["aggregate"]["resultado"] == "3234.75"
    assert data["schema"]["Valor"]["type"] == "decimal"
    assert data["nulls_excluded"] == 1
    assert data["row_count"] == 4


def test_medium_risk_does_not_select_risk_state_and_typo_does_not_count_all(corpus):
    path = source(corpus, "RiskyUsers.csv", "Usuário;Estado do risco;Nível de risco;Status\nA;Em risco;Alto;Ativo\nB;Em risco;Médio;Ativo\nC;Comprometimento confirmado;Médio;Inativo\n")
    for question in ("Quantos usuários estão em risco médio?", "Quantos tem risco avarege?"):
        result = analyze_structured_question(corpus, "infraestrutura", question)
        assert result.matched_count == 2
        assert result.filter_column == "Nível de risco"
    assert analyze_structured_question(corpus, "infraestrutura", "Quantos tem risco inexplicavel?").operation == "clarification"
    assert analyze_structured_question(corpus, "infraestrutura", "Quantos usuários com risco médio e Status Ativo?").matched_count == 1


def test_correlation_is_not_causation(corpus):
    source(corpus, "metricas.csv", "Carga;Latencia\n1;10\n2;20\n3;30\n4;40\n")
    result = ask(corpus, "Qual a correlação entre Carga e Latencia em metricas?")
    assert result.context_package["structured_data"]["aggregate"]["resultado"] == 1.0
    assert "não demonstra causalidade" in result.answer


def test_mixed_pdf_uses_real_ocr_and_retains_page_provenance(corpus):
    path = corpus / "infraestrutura" / "escaneado.pdf"
    scanned_pdf(path, "MANUAL ATLAS\nA verificacao do backup ocorre diariamente.\nA retencao dos arquivos e de 45 dias.\nO responsavel deve registrar cada verificacao.", mixed=True)
    pages = extract_pages(path)
    assert pages_ready(pages), pages
    assert [p.method for p in pages] == ["native", "ocr"]
    assert "45 dias" in pages[1].text
    prepare(corpus, path)
    result = retrieve(corpus, "infraestrutura", "Qual a retenção no arquivo escaneado?", policy_for("infraestrutura"))
    assert any(e.chunk.page == 2 and "45 dias" in e.chunk.text for e in result.evidence)
    record = pipeline_documents(corpus, "infraestrutura")[0]
    assert record["status"] == "READY"
    assert any(e["stage"] == "OCR" for e in record["events"])


def test_bad_ocr_page_blocks_entire_document(corpus):
    path = corpus / "infraestrutura" / "ilegivel.pdf"
    scanned_pdf(path, "", mixed=True)
    pages = extract_pages(path)
    assert not pages_ready(pages)
    result = record_document_pipeline(corpus, "infraestrutura", path)
    assert result["status"] == "QUARANTINED"
    assert not ingest_module(corpus, "infraestrutura")


def test_retrieval_never_runs_heavy_extraction(corpus, monkeypatch):
    path = corpus / "infraestrutura" / "novo.pdf"
    scanned_pdf(path, "Atlas exige backup diario.")
    def forbidden(*args, **kwargs):
        raise AssertionError("OCR ran during query")
    monkeypatch.setattr("api.document_pages.extract_pages", forbidden)
    result = retrieve(corpus, "infraestrutura", "Resuma novo.pdf", policy_for("infraestrutura"))
    assert not result.has_quality_evidence


def test_no_evidence_and_unsupported_numeric_claim(corpus):
    path = source(corpus, "politica.txt", "O prazo de retenção dos logs do Atlas é de 30 dias. O prazo de retenção do Beta é de 90 dias.")
    prepare(corpus, path)
    result = ask(corpus, "Segundo o documento politica, qual a senha do roteador?")
    assert not result.evidence_found
    assert "não" in result.answer.lower()
    evidence = retrieve(corpus, "infraestrutura", "Resuma politica", policy_for("infraestrutura"))
    assert verify_answer("O prazo de retenção dos logs do Atlas é de 90 dias. [E1]", evidence, "infraestrutura").status == "rejected"


def test_management_loss_question_requires_area_financial_data(corpus):
    source(corpus, "governanca.txt", "A governança organiza responsabilidades, controles e prestação de contas.", "gestao-empresarial")
    result = ask(corpus, "É possível identificar que áreas podem estar dando prejuízo para organização?", "gestao-empresarial")
    assert result.provider == "policy"
    assert result.verification_status == "verified"
    assert "não é possível apontar quais áreas" in result.answer
    assert "receitas" in result.answer and "custos" in result.answer and "resultado líquido" in result.answer
    assert "seria uma hipótese" in result.answer


def test_management_production_variance_rejects_document_production_false_positive(corpus):
    source(
        corpus,
        "governanca-documental.txt",
        "A produção diária de documentos deve seguir o fluxo de tramitação. "
        "O sistema organiza processos, metas institucionais e indicadores gerais.",
        "gestao-empresarial",
    )
    result = ask(corpus, "Como vejo desvios de produção?", "gestao-empresarial")
    assert not result.evidence_found
    assert result.provider == "policy"
    assert "desvios de produção" in result.answer
    assert "produção diária de documentos" not in result.answer
    assert "meta/planejado" in result.answer


def test_management_production_variance_accepts_related_operational_evidence(corpus):
    path = source(
        corpus,
        "relatorio-producao.txt",
        "Relatório de produção. A meta planejada era 100 unidades e o realizado "
        "foi 80. O desvio de produção foi de -20%, com perdas de 5 unidades.",
        "gestao-empresarial",
    )
    prepare(corpus, path)
    result = ask(corpus, "Como vejo desvios de produção?", "gestao-empresarial")
    assert result.evidence_found
    assert result.provider == "local-rag"
    assert result.sources == [path.name]
    assert "meta planejada" in result.answer
    assert "desvio de produção" in result.answer


def test_pipeline_diagnostics_are_restricted_to_the_designated_admin():
    require_admin({"sub": "AG000001"})
    for user_code in ("IN000001", "RH000001", "AG000002", ""):
        with pytest.raises(HTTPException) as error:
            require_admin({"sub": user_code})
        assert error.value.status_code == 403


def test_medical_definition_excludes_codes_and_external_markdown_is_preserved(corpus):
    table = source(corpus, "CID-10.csv", "Codigo;Descricao\nJ09;Influenza gripe aviaria\nJ100;Influenza com pneumonia\n", "medicina")
    text = source(corpus, "guia.txt", "Gripe é uma infecção respiratória causada por vírus influenza. O material apresenta informações gerais sobre prevenção.", "medicina")
    prepare(corpus, table, text)
    result = ask(corpus, "Defina gripe?", "medicina")
    assert "infecção respiratória" in result.answer
    assert "J09" not in result.answer
    assert table.name not in result.sources
    formatted = _format_external_assist_answer("**Admissão**\n\nOrganiza documentos.\n\n**Férias**\n\nControla períodos.", "openai", "O que DP faz?", "pt-BR", "structured")
    assert "**Admissão**" in formatted and "\n\n" in formatted
    assert "Base documental" not in formatted


def test_trigger_procedure_is_portuguese_and_not_navigation(corpus):
    path = source(corpus, "www-zabbix-com-guide.md", "Copyright notice On this page Overview Testing expressions To configure a trigger, do the following: Go to: Data collection > Hosts. Click Triggers in the row of the host. Create trigger to the right (or on the trigger name to edit an existing trigger). Enter parameters of the trigger in the form.")
    prepare(corpus, path)
    result = ask(corpus, "Como que defino um trigger?")
    assert "Acesse" in result.answer and "clique" in result.answer
    assert "Copyright" not in result.answer and "Base documental\n- The" not in result.answer
